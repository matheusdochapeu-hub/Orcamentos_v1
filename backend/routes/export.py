"""
Rotas para exportação de resultados em Excel e PDF.
"""
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
import io
import json
import re
from datetime import datetime

router = APIRouter(prefix="/api/export", tags=["Exportação"])


def _sanitizar_nome_arquivo(nome: str) -> str:
    """Sanitiza o nome do arquivo para o cabeçalho Content-Disposition."""
    return re.sub(r'[^a-zA-Z0-9_.-]', '_', nome)


class ExportRequest(BaseModel):
    # Aceita tanto os nomes antigos quanto os enviados pelo frontend.
    nome_busca: Optional[str] = ""
    catdescricao: Optional[str] = ""
    catcodigo: Optional[str] = ""
    tipo: Optional[str] = ""
    cidade: Optional[str] = ""
    raio_km: Optional[int] = 0
    tipo_busca: Optional[str] = ""
    empresas: list[dict] = []

    # Aliases enviados pelo frontend (app.js)
    titulo: Optional[str] = ""
    codigo_cat: Optional[str] = ""
    tipo_cat: Optional[str] = ""
    fontes_usadas: Optional[list[str]] = None

    def normalizar(self) -> "ExportRequest":
        """Preenche os campos canônicos a partir dos aliases do frontend."""
        self.catdescricao = self.catdescricao or self.titulo or ""
        self.catcodigo = self.catcodigo or self.codigo_cat or ""
        self.tipo = self.tipo or self.tipo_cat or ""
        self.tipo_busca = self.tipo_busca or "—"
        self.nome_busca = self.nome_busca or self.catdescricao
        return self


# ─── EXCEL ────────────────────────────────────────────────────────────────────

@router.post("/excel")
async def exportar_excel(req: ExportRequest):
    """Gera planilha Excel com a lista de empresas."""
    req = req.normalizar()
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "Biblioteca openpyxl não instalada")

    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"

    # Paleta de cores
    azul_escuro = "1A3A5C"
    azul_claro  = "D6E4F0"
    verde       = "27AE60"
    amarelo     = "F39C12"
    vermelho    = "E74C3C"
    cinza       = "F2F2F2"

    # ── Cabeçalho da busca ──────────────────────────────────────────────────
    ws.merge_cells("A1:K1")
    ws["A1"] = "🔍 RELATÓRIO DE BUSCA DE FORNECEDORES"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    info_linhas = [
        ("Item / Serviço:", req.catdescricao),
        (f"Código {req.tipo}:", req.catcodigo),
        ("Cidade:", f"{req.cidade} — Raio: {req.raio_km} km"),
        ("Tipo de busca:", req.tipo_busca.capitalize()),
        ("Data/Hora:", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Total de empresas:", str(len(req.empresas))),
    ]
    for linha_num, (label, valor) in enumerate(info_linhas, start=2):
        ws.merge_cells(f"A{linha_num}:C{linha_num}")
        ws[f"A{linha_num}"] = label
        ws[f"A{linha_num}"].font = Font(bold=True)
        ws[f"A{linha_num}"].fill = PatternFill("solid", fgColor=azul_claro)
        ws.merge_cells(f"D{linha_num}:K{linha_num}")
        ws[f"D{linha_num}"] = valor

    # ── Cabeçalho da tabela ─────────────────────────────────────────────────
    linha_header = len(info_linhas) + 3
    headers = ["#", "Relevância", "Nome", "CNPJ", "CNAE", "Endereço", "Telefone", "WhatsApp", "E-mail", "Site", "Avaliação Google"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=linha_header, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=azul_escuro)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            bottom=Side(style="medium", color="FFFFFF"),
            right=Side(style="thin", color="FFFFFF"),
        )
    ws.row_dimensions[linha_header].height = 22

    # ── Dados das empresas ──────────────────────────────────────────────────
    for i, emp in enumerate(req.empresas):
        row = linha_header + 1 + i
        rel = emp.get("relevancia", 0)

        # Cor da linha baseada na relevância
        if rel >= 70:
            cor_linha = "D5F5E3"  # verde claro
        elif rel >= 40:
            cor_linha = "FEF9E7"  # amarelo claro
        else:
            cor_linha = "FDEDEC"  # vermelho claro

        valores = [
            i + 1,
            f"{rel}%",
            emp.get("nome", "") or emp.get("razao_social", ""),
            emp.get("cnpj", "") or "—",
            (emp.get("cnae_principal", "") or "—"),
            emp.get("endereco", ""),
            emp.get("telefone", ""),
            emp.get("whatsapp", ""),
            emp.get("email", ""),
            emp.get("site", ""),
            f"{emp.get('avaliacao_google') or emp.get('avaliacao', 0):.1f} ⭐" if emp.get("avaliacao_google") else "—",
        ]
        for col, val in enumerate(valores, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = PatternFill("solid", fgColor=cor_linha if i % 2 == 0 else cinza)
            cell.alignment = Alignment(vertical="center", wrap_text=(col in [3, 4]))
            cell.border = Border(
                bottom=Side(style="thin", color="CCCCCC"),
                right=Side(style="thin", color="CCCCCC"),
            )

        ws.row_dimensions[row].height = 20

    # ── Largura das colunas ─────────────────────────────────────────────────
    larguras = [5, 11, 32, 20, 14, 42, 17, 17, 28, 36, 16]
    for col, larg in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(col)].width = larg

    # ── Rodapé ─────────────────────────────────────────────────────────────
    ultima_linha = linha_header + len(req.empresas) + 2
    ws.merge_cells(f"A{ultima_linha}:K{ultima_linha}")
    ws[f"A{ultima_linha}"] = "Gerado por Sistema de Orçamentos Públicos"
    ws[f"A{ultima_linha}"].font = Font(italic=True, color="888888")
    ws[f"A{ultima_linha}"].alignment = Alignment(horizontal="center")

    # Salva em memória
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nome_base = f"orcamento_{req.catcodigo}_{req.cidade}_{datetime.now().strftime('%Y%m%d_%H%M')}"
    nome_arquivo = f"{_sanitizar_nome_arquivo(nome_base)}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome_arquivo}"'},
    )


# ─── PDF ──────────────────────────────────────────────────────────────────────

@router.post("/pdf")
async def exportar_pdf(req: ExportRequest):
    """Gera PDF com a lista de empresas."""
    req = req.normalizar()
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph,
            Spacer, HRFlowable,
        )
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
    except ImportError:
        raise HTTPException(500, "Biblioteca reportlab não instalada")

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    azul = colors.HexColor("#1A3A5C")
    verde = colors.HexColor("#27AE60")
    amarelo = colors.HexColor("#F39C12")
    vermelho = colors.HexColor("#E74C3C")

    estilo_titulo = ParagraphStyle(
        "Titulo",
        parent=styles["Title"],
        textColor=azul,
        fontSize=16,
        spaceAfter=4,
    )
    estilo_sub = ParagraphStyle(
        "Sub",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.gray,
        spaceAfter=2,
    )
    estilo_cell = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
    )

    elementos = []

    # Título
    elementos.append(Paragraph("🔍 Relatório de Busca de Fornecedores", estilo_titulo))
    elementos.append(HRFlowable(width="100%", thickness=2, color=azul))
    elementos.append(Spacer(1, 0.3 * cm))

    # Informações da busca
    info_text = (
        f"<b>Item:</b> {req.catdescricao} &nbsp;|&nbsp; "
        f"<b>Código {req.tipo}:</b> {req.catcodigo} &nbsp;|&nbsp; "
        f"<b>Cidade:</b> {req.cidade} &nbsp;|&nbsp; "
        f"<b>Raio:</b> {req.raio_km} km &nbsp;|&nbsp; "
        f"<b>Busca:</b> {req.tipo_busca.capitalize()} &nbsp;|&nbsp; "
        f"<b>Total:</b> {len(req.empresas)} empresa(s) &nbsp;|&nbsp; "
        f"<b>Data:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )
    elementos.append(Paragraph(info_text, estilo_sub))
    elementos.append(Spacer(1, 0.4 * cm))

    # Tabela
    header_row = ["#", "Relev.", "Nome", "CNPJ", "Endereço", "Telefone", "E-mail", "Site", "⭐"]
    col_widths = [0.8*cm, 1.3*cm, 5*cm, 3.2*cm, 5.5*cm, 2.8*cm, 4.2*cm, 4.2*cm, 1.2*cm]

    dados = [header_row]
    estilos_linhas = []

    for i, emp in enumerate(req.empresas):
        rel = emp.get("relevancia", 0)
        if rel >= 70:
            cor = colors.HexColor("#D5F5E3")
        elif rel >= 40:
            cor = colors.HexColor("#FEF9E7")
        else:
            cor = colors.HexColor("#FDEDEC")

        estilos_linhas.append(("BACKGROUND", (0, i + 1), (-1, i + 1), cor))

        row = [
            str(i + 1),
            f"{rel}%",
            Paragraph(emp.get("nome", "") or emp.get("razao_social", ""), estilo_cell),
            emp.get("cnpj", "") or "—",
            Paragraph(emp.get("endereco", ""), estilo_cell),
            emp.get("telefone", "") or "—",
            Paragraph(emp.get("email", "") or "—", estilo_cell),
            Paragraph(emp.get("site", "") or "—", estilo_cell),
            f"{emp.get('avaliacao_google') or 0:.1f}" if emp.get("avaliacao_google") else "—",
        ]
        dados.append(row)

    tabela = Table(dados, colWidths=col_widths, repeatRows=1)
    tabela.setStyle(TableStyle([
        # Header
        ("BACKGROUND", (0, 0), (-1, 0), azul),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        # Dados
        ("FONTSIZE", (0, 1), (-1, -1), 7),
        ("ALIGN", (0, 1), (1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ] + estilos_linhas))

    elementos.append(tabela)
    elementos.append(Spacer(1, 0.5 * cm))
    elementos.append(Paragraph(
        f"<i>Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')} — Sistema de Orçamentos Públicos</i>",
        estilo_sub,
    ))

    doc.build(elementos)
    output.seek(0)

    nome_base = f"orcamento_{req.catcodigo}_{req.cidade}_{datetime.now().strftime('%Y%m%d_%H%M')}"
    nome_arquivo = f"{_sanitizar_nome_arquivo(nome_base)}.pdf"
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{nome_arquivo}"'},
    )
